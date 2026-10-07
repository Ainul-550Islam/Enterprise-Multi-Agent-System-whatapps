"""Bounded parsers for untrusted RAG/Data Hub content.

Parsing turns source bytes into plain text. It is also the first place
untrusted content enters the system, so scripts, styles, executable markup,
malformed containers, decompression bombs, and excessive page/row counts are
rejected or removed before prompt construction. Rich-format parsers are
optional dependencies and are imported only when their format is requested.
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
import zipfile
from typing import Final

from mas_orchestrator.contracts.errors import ConfigurationError, InvalidRequestError
from mas_orchestrator.rag.models import DocumentFormat

MAX_PARSED_CHARS: Final[int] = 8_000_000
MAX_PDF_PAGES: Final[int] = 500
MAX_CSV_ROWS: Final[int] = 100_000
MAX_CSV_COLUMNS: Final[int] = 2_000
MAX_DOCX_MEMBERS: Final[int] = 5_000
MAX_DOCX_UNCOMPRESSED_BYTES: Final[int] = 64 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO: Final[int] = 200

# --- HTML handling ---------------------------------------------------------

_SCRIPT: Final[re.Pattern[str]] = re.compile(
    r"<(script|style|noscript|template)\b.*?</\1\s*>", re.IGNORECASE | re.DOTALL
)
_COMMENT: Final[re.Pattern[str]] = re.compile(r"<!--.*?-->", re.DOTALL)
_TAG: Final[re.Pattern[str]] = re.compile(r"<[^>]+>")
_BLOCK_TAGS: Final[re.Pattern[str]] = re.compile(
    r"</?(p|div|br|li|tr|h[1-6]|section|article)\s*/?>", re.IGNORECASE
)
_ENTITY_REF: Final[re.Pattern[str]] = re.compile(r"&[a-zA-Z#][a-zA-Z0-9]{0,8};")
_WS: Final[re.Pattern[str]] = re.compile(r"[ \t]+")
_BLANKS: Final[re.Pattern[str]] = re.compile(r"\n{3,}")

# --- Markdown handling -----------------------------------------------------

_MD_CODE: Final[re.Pattern[str]] = re.compile(r"```.*?```", re.DOTALL)
_MD_LINK: Final[re.Pattern[str]] = re.compile(r"\[([^\]]*)\]\([^)]*\)")
_MD_IMAGE: Final[re.Pattern[str]] = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_MD_HEADING: Final[re.Pattern[str]] = re.compile(r"^#{1,6}\s*", re.MULTILINE)
_MD_EMPHASIS: Final[re.Pattern[str]] = re.compile(r"(\*\*|__|\*|_|~~|`)")
_MD_HTML: Final[re.Pattern[str]] = re.compile(r"<[^>]+>")


class Parser:
    """Turns bounded document content into normalized plain text."""

    def parse(self, content: str | bytes, format: DocumentFormat) -> str:  # noqa: A002
        """Parses one document according to its trusted, declared format.

        Args:
            content: Raw text or source bytes, already bounded by the caller.
            format: Format derived from the source MIME type, not model output.

        Returns:
            Plain text, never more than :data:`MAX_PARSED_CHARS` characters.

        Raises:
            ConfigurationError: A requested optional parser is not installed.
            InvalidRequestError: The bytes are malformed or exceed parser caps.
        """
        if isinstance(content, bytes):
            if not content:
                return ""
        elif not content.strip():
            return ""

        if format is DocumentFormat.PDF:
            parsed = _parse_pdf(content)
        elif format is DocumentFormat.DOCX:
            parsed = _parse_docx(content)
        else:
            text = _decode_text(content)
            dispatch = {
                DocumentFormat.TEXT: _parse_text,
                DocumentFormat.MARKDOWN: _parse_markdown,
                DocumentFormat.JSON: _parse_json,
                DocumentFormat.HTML: _parse_html,
                DocumentFormat.CSV: _parse_csv,
            }
            try:
                parser = dispatch[format]
            except KeyError as error:
                raise InvalidRequestError(
                    "document format is not supported",
                    details={"format": str(format)},
                ) from error
            parsed = parser(text)

        if len(parsed) > MAX_PARSED_CHARS:
            raise InvalidRequestError(
                "parsed document exceeds the maximum size",
                details={"limit": MAX_PARSED_CHARS, "size": len(parsed)},
            )
        return parsed


def _decode_text(content: str | bytes) -> str:
    """Decodes text formats strictly; replacement characters hide corruption."""
    if isinstance(content, str):
        return content
    try:
        return content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise InvalidRequestError(
            "text document is not valid UTF-8",
            details={"byte_offset": error.start},
        ) from error


def _parse_text(content: str) -> str:
    """Normalises whitespace only."""
    return _normalise(content)


def _parse_markdown(content: str) -> str:
    """Reduces Markdown to readable text while retaining code and labels."""
    text = _MD_CODE.sub(lambda match: match.group(0).strip("`"), content)
    text = _MD_IMAGE.sub("", text)
    text = _MD_LINK.sub(r"\1", text)
    text = _MD_HEADING.sub("", text)
    text = _MD_EMPHASIS.sub("", text)
    text = _MD_HTML.sub(" ", text)
    return _normalise(text)


def _parse_json(content: str) -> str:
    """Flattens JSON into labelled lines, preserving object paths."""
    try:
        payload = json.loads(content)
    except json.JSONDecodeError as error:
        raise InvalidRequestError(
            "document is not valid JSON", details={"position": error.pos}
        ) from error

    lines: list[str] = []

    def walk(value: object, path: str) -> None:
        if isinstance(value, dict):
            for key, item in value.items():
                walk(item, f"{path}.{key}" if path else str(key))
        elif isinstance(value, list):
            for index, item in enumerate(value):
                walk(item, f"{path}[{index}]")
        else:
            lines.append(f"{path}: {value}")

    walk(payload, "")
    return _normalise("\n".join(lines))


def _parse_csv(content: str) -> str:
    """Flattens bounded CSV into labelled cells without spreadsheet execution."""
    try:
        reader = csv.reader(io.StringIO(content, newline=""), strict=True)
        rows: list[list[str]] = []
        for row_number, row in enumerate(reader, start=1):
            if row_number > MAX_CSV_ROWS:
                raise InvalidRequestError(
                    "CSV exceeds the maximum row count",
                    details={"limit": MAX_CSV_ROWS},
                )
            if len(row) > MAX_CSV_COLUMNS:
                raise InvalidRequestError(
                    "CSV row exceeds the maximum column count",
                    details={"limit": MAX_CSV_COLUMNS, "row": row_number},
                )
            rows.append(row)
    except csv.Error as error:
        raise InvalidRequestError("document is not valid CSV") from error
    if not rows:
        return ""
    header = rows[0]
    lines: list[str] = []
    for row_number, row in enumerate(rows[1:], start=2):
        for column, value in enumerate(row):
            name = header[column].strip() if column < len(header) else f"column_{column + 1}"
            if value.strip():
                lines.append(f"row {row_number} {name}: {value}")
    return _normalise("\n".join(lines))


def _parse_html(content: str) -> str:
    """Removes executable/presentational HTML before text extraction."""
    text = _SCRIPT.sub(" ", content)
    text = _COMMENT.sub(" ", text)
    text = _BLOCK_TAGS.sub("\n", text)
    text = _TAG.sub(" ", text)
    text = html.unescape(text)
    text = _ENTITY_REF.sub(" ", text)
    return _normalise(text)


def _parse_pdf(content: str | bytes) -> str:
    """Extracts bounded, text-only PDF pages using the optional pypdf parser."""
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ConfigurationError(
            "PDF ingestion requires the optional ingest extra",
            details={"extra": "ingest", "parser": "pypdf"},
        ) from error
    raw = content if isinstance(content, bytes) else content.encode("utf-8")
    try:
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise InvalidRequestError(
                "PDF exceeds the maximum page count",
                details={"limit": MAX_PDF_PAGES, "pages": len(reader.pages)},
            )
        pages: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            if len("\n".join(pages)) > MAX_PARSED_CHARS:
                break
            pages.append(f"Page {page_number}\n{page.extract_text() or ''}")
        return _normalise("\n\n".join(pages))
    except InvalidRequestError:
        raise
    except Exception as error:
        raise InvalidRequestError("PDF document is malformed or unreadable") from error


def _parse_docx(content: str | bytes) -> str:
    """Extracts DOCX paragraphs/tables after ZIP-bomb and member validation."""
    try:
        from docx import Document
    except ImportError as error:
        raise ConfigurationError(
            "DOCX ingestion requires the optional ingest extra",
            details={"extra": "ingest", "parser": "python-docx"},
        ) from error
    raw = content if isinstance(content, bytes) else content.encode("utf-8")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            total_size = sum(member.file_size for member in members)
            if len(members) > MAX_DOCX_MEMBERS or total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise InvalidRequestError(
                    "DOCX package exceeds parser safety limits",
                    details={
                        "member_limit": MAX_DOCX_MEMBERS,
                        "uncompressed_byte_limit": MAX_DOCX_UNCOMPRESSED_BYTES,
                    },
                )
            for member in members:
                if member.file_size > 0 and member.compress_size == 0:
                    raise InvalidRequestError("DOCX package has an invalid compressed member")
                if (
                    member.compress_size > 0
                    and member.file_size / member.compress_size > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise InvalidRequestError("DOCX package exceeds the compression-ratio limit")
            if "word/document.xml" not in archive.namelist():
                raise InvalidRequestError("DOCX package has no Word document part")
        document = Document(io.BytesIO(raw))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return _normalise("\n".join(lines))
    except InvalidRequestError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as error:
        raise InvalidRequestError("DOCX document is malformed or unreadable") from error


def _normalise(text: str) -> str:
    """Collapses whitespace into a stable, citation-friendly form."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS.sub(" ", text)
    text = _BLANKS.sub("\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def detect_format(
    content: str | bytes, *, declared: DocumentFormat | None = None
) -> DocumentFormat:
    """Guesses text/PDF format when none is declared; trusted MIME wins."""
    if declared is not None:
        return declared
    if isinstance(content, bytes):
        if content.startswith(b"%PDF-"):
            return DocumentFormat.PDF
        prefix = content[:4096].lstrip()
        try:
            stripped = prefix.decode("utf-8-sig", errors="strict")
        except UnicodeDecodeError:
            return DocumentFormat.TEXT
    else:
        stripped = content.lstrip()
    if stripped.startswith("<"):
        return DocumentFormat.HTML
    if stripped.startswith("{") or stripped.startswith("["):
        return DocumentFormat.JSON
    if _MD_HEADING.search(stripped) or "```" in stripped or _MD_LINK.search(stripped):
        return DocumentFormat.MARKDOWN
    if "," in stripped.splitlines()[0] if stripped.splitlines() else False:
        return DocumentFormat.CSV
    return DocumentFormat.TEXT
